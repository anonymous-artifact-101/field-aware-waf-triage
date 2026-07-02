
from __future__ import annotations

import hashlib
import os
from typing import Dict, Iterator, List, Optional, Sequence, Tuple

from src.data.parsers.apache_combined import _parse_request_line

__all__ = [
    "parse_file",
    "parse_row",
    "RECORD_KEYS",
    "MODEL_FIELDS",
    "SHEET_NAME",
    "HEADER",
]

MODEL_FIELDS: Tuple[str, ...] = ("method", "path", "query", "ua", "status", "timing")

RECORD_KEYS: Tuple[str, ...] = (
    "unique_id",
    "day",
    "method",
    "path",
    "query",
    "ua",
    "status",
    "timing",
    "bytes",
    "referer",
    "ip_hash",
    "country",
    "detected_raw",
    "label",
    "attack_subtype",
)

SHEET_NAME = "Main"

HEADER: Tuple[str, ...] = (
    "ip", "datetime", "gmt", "request", "status", "size",
    "referer", "browser", "country", "detected",
)

_DAY = "Jul-2019"

_IP_SALT = "pecti-apache-indo-v1"

def _to_int(token: object) -> int:
    if token is None:
        return 0
    s = str(token).strip()
    if not s or s == "-":
        return 0
    try:
        return int(float(s))
    except ValueError:
        return 0

def _hash_ip(ip: object) -> str:
    s = "" if ip is None else str(ip).strip()
    if not s:
        return ""
    return hashlib.sha256((_IP_SALT + s).encode("utf-8")).hexdigest()[:12]

def _label_from_detected(detected: object) -> str:
    v = ("" if detected is None else str(detected)).strip().upper()
    return "benign" if v == "AMAN" else "flagged"

def parse_row(cells: Sequence[object], row_index: int = 0) -> Optional[Dict]:
    if cells is None or len(cells) < len(HEADER):
        return None

    ip, _datetime, _gmt, request, status, size, referer, browser, country, detected = \
        cells[:len(HEADER)]

    method, path, query = _parse_request_line("" if request is None else str(request))
    if method == "":
        return None

    ref = "" if referer is None else str(referer).strip()
    return {
        "unique_id": f"apacheindo-{row_index:09d}",
        "day": _DAY,
        "method": method,
        "path": path,
        "query": query,
        "ua": "" if browser is None else str(browser).strip(),
        "status": _to_int(status),
        "timing": 0,
        "bytes": _to_int(size),
        "referer": "" if ref == "-" else ref,
        "ip_hash": _hash_ip(ip),
        "country": "" if country is None else str(country).strip(),
        "detected_raw": "" if detected is None else str(detected).strip(),
        "label": _label_from_detected(detected),
        "attack_subtype": "",
    }

def parse_file(path: "str | os.PathLike[str]") -> Iterator[Dict]:
    try:
        import openpyxl
    except ImportError as exc:
        raise SystemExit(
            "[apache_indo] openpyxl is required to read the .xlsx dataset; "
            "pip install openpyxl"
        ) from exc

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(
            f"[apache_indo] sheet {SHEET_NAME!r} not found in {path}; "
            f"sheets present: {wb.sheetnames}"
        )
    ws = wb[SHEET_NAME]
    data_index = 0
    for r_i, row in enumerate(ws.iter_rows(values_only=True)):
        if r_i == 0:
            continue
        record = parse_row(row, row_index=data_index)
        if record is not None:
            yield record
            data_index += 1

def _self_test() -> None:
    import json

    sample_rows: List[Sequence[object]] = [
        ["36.90.2.19", "2019-07-23 09:30:17", "+0700]", "GET /bkd_baru/ HTTP/1.1",
         "200", "5198", "http://sc.syekhnurjati.ac.id/bkd/", "Mozilla/5.0",
         "Indonesia", "AMAN"],
        ["1.2.3.4", "2019-07-23 09:31:00", "+0700]",
         "GET /x/nyil.php?path=C:/xampp HTTP/1.1", "200", "12", "-",
         "curl/7.64", "Russia", "BAHAYA"],
    ]
    for i, row in enumerate(sample_rows):
        rec = parse_row(row, row_index=i)
        print(json.dumps(rec, sort_keys=True, ensure_ascii=False))

if __name__ == "__main__":
    _self_test()
